import os
from docx import Document
from openpyxl import Workbook
import pymupdf as fitz

class DocumentManager:
    def __init__(self, output_dir: str = "generated_docs"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def read_any_file(self, file_path: str) -> dict:
        """Reads and extracts text from PDFs, Word docs, text files, code, etc."""
        # Clean quotes
        clean_path = file_path.strip("\"' ")
        if not os.path.exists(clean_path):
            raise FileNotFoundError(f"File '{clean_path}' was not found on your system.")

        filename = os.path.basename(clean_path)
        _, ext = os.path.splitext(filename)
        ext = ext.lower()
        extracted_text = ""

        if ext == ".pdf":
            doc = fitz.open(clean_path)
            for page in doc:
                extracted_text += page.get_text() + "\n"
        elif ext in [".docx", ".doc"]:
            doc = Document(clean_path)
            extracted_text = "\n".join([p.text for p in doc.paragraphs if p.text])
        elif ext in [".txt", ".csv", ".json", ".py", ".md", ".log", ".yaml", ".yml", ".ini"]:
            with open(clean_path, "r", encoding="utf-8", errors="ignore") as f:
                extracted_text = f.read()
        else:
            extracted_text = f"Binary/Media file: {filename} (Size: {os.path.getsize(clean_path)} bytes)"

        return {
            "filename": filename,
            "filepath": clean_path,
            "extension": ext,
            "text": extracted_text.strip(),
            "size_kb": round(os.path.getsize(clean_path) / 1024, 2)
        }

    def create_word_doc(self, title: str, paragraphs: list, filename: str = "document.docx") -> str:
        doc = Document()
        doc.add_heading(title, 0)
        for p in paragraphs:
            doc.add_paragraph(p)
        filepath = os.path.join(self.output_dir, filename)
        doc.save(filepath)
        return filepath

    def create_excel_sheet(self, headers: list, rows: list, filename: str = "report.xlsx") -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "Diablo Report"
        ws.append(headers)
        for row in rows:
            ws.append(row)
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)
        return filepath